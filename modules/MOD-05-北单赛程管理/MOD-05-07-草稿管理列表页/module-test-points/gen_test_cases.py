# -*- coding: utf-8 -*-
"""按模板格式生成 MOD-05-07 草稿管理列表页 第13轮冲刺（PBL-618）测试用例 Excel"""
import os
import shutil
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

BASE_DIR = r"D:\pythonProject\test-league-platform-ai"
MOD_DIR = os.path.join(BASE_DIR, r"modules\MOD-05-北单赛程管理\MOD-05-07-草稿管理列表页")
OUT_DIR = os.path.join(MOD_DIR, "module-test-cases")
os.makedirs(OUT_DIR, exist_ok=True)
OUT = os.path.join(OUT_DIR, "测试用例-草稿管理列表页-第13轮冲刺.xlsx")
GLOBAL_OUT = os.path.join(BASE_DIR, "test-cases", "测试用例-草稿管理列表页-第13轮冲刺.xlsx")

# 列: A序号 B冲刺 C测试项 D测试点 E前置条件 F执行步骤 G预期结果 H测试结果1 I测试结果2 J测试结果3
cases = [
    ("按钮显隐", "球队未匹配纳米数据源时，快速关联球队弹窗展示【匹配纳米数据源】按钮",
     "草稿列表中存在未关联赛季的球队，且 m_team.nami_id_match_status=0（如 s_match=116316 主队 7328 汉威尔镇）",
     "1.点击该球队名称旁的未关联标识",
     "弹出快速关联球队弹窗，展示【匹配纳米数据源】按钮"),

    ("", "球队已匹配纳米数据源时，快速关联球队弹窗不展示【匹配纳米数据源】按钮",
     "草稿列表中存在未关联赛季的球队，且 m_team.nami_id_match_status=1（如 s_match=116314 客队 1309 俄斯特拉发）",
     "1.点击该球队名称旁的未关联标识",
     "弹出快速关联球队弹窗，不展示【匹配纳米数据源】按钮"),

    ("", "主队、客队分别打开弹窗，按钮按各自球队的纳米匹配状态显隐",
     "草稿中同一场比赛主队未匹配纳米（status=0）、客队已匹配纳米（status=1），且两队均未关联当前赛季（如 s_match=116314：主队 6798 克拉德诺 status=0，客队 1309 俄斯特拉发 status=1）",
     "1.点击主队未关联标识查看弹窗\n2.关闭后点击客队未关联标识查看弹窗",
     "1.主队弹窗展示【匹配纳米数据源】按钮\n2.客队弹窗不展示【匹配纳米数据源】按钮"),

    ("跳转匹配球队页", "点击【匹配纳米数据源】，在新标签页打开球队管理中当前球队的匹配球队页",
     "已打开未匹配纳米球队的快速关联球队弹窗",
     "1.点击弹窗中的【匹配纳米数据源】按钮",
     "浏览器在新标签页打开球队管理中当前球队的匹配球队页"),

    ("", "匹配球队页默认显示\"纳米数据\"标签页",
     "同用例 4，新标签页已打开匹配球队页",
     "1.查看匹配球队页当前选中的数据源标签页",
     "默认选中并显示\"纳米数据\"标签页"),

    ("", "匹配球队页球队名称筛选默认填入当前选中球队的球队简称",
     "同用例 4，当前选中球队在 m_team 表的球队简称为\"汉威尔镇\"",
     "1.查看匹配球队页顶部的球队名称筛选输入框",
     "球队名称筛选框默认填入\"汉威尔镇\"（当前选中球队的 short_name_zh）"),

    ("", "球队简称为纯英文或含特殊字符时，匹配球队页球队名称筛选框按原样完整带入展示",
     "草稿中存在未关联赛季且未匹配纳米的纯英文简称球队（如 7149 'FC Murom'）或含特殊字符简称球队（如 7512 '拉姆/索尔兹伯里'）",
     "1.点击该球队未关联标识打开弹窗\n2.点击【匹配纳米数据源】打开新标签页\n3.查看匹配球队页球队名称筛选框",
     "球队名称筛选框按页面/数据库字段存的原样完整带入展示（'FC Murom' 或 '拉姆/索尔兹伯里'），无截断或丢失"),

    ("弹窗状态与关联流程", "新打开标签页后，原草稿页快速关联球队弹窗不关闭、继续保留，已填写的北单简称、北单全称、命名类型内容不变",
     "已打开未匹配纳米球队的快速关联弹窗，并填写了北单简称、北单全称、命名类型",
     "1.点击【匹配纳米数据源】打开新标签页\n2.切回原草稿管理列表页查看弹窗",
     "原草稿页快速关联球队弹窗不关闭、继续保留，已填写的北单简称、北单全称、命名类型内容不变"),

    ("", "在新标签页完成纳米数据源匹配后，未刷新草稿页时，原草稿页弹窗不会有任何改变，仍停留在原状态",
     "点击【匹配纳米数据源】打开新标签页，并在新标签页勾选纳米球队完成匹配保存",
     "1.直接切回未刷新的原草稿管理列表页查看快速关联弹窗",
     "由于草稿页未实时刷新，原草稿页弹窗不会有任何改变，仍停留在原状态（弹窗不关闭、【匹配纳米数据源】按钮仍显示、已填内容不变）"),

    ("", "点击【匹配纳米数据源】（或在新标签页完成匹配）后，切回原草稿页弹窗点击【确定】仍可正常完成球队关联赛季",
     "同用例 8 或用例 9，切回原草稿管理列表页弹窗，填写的北单简称、北单全称合规且唯一",
     "1.在快速关联弹窗点击【确定】",
     "提示关联成功，关闭弹窗，球队新增至当前比赛赛季下并更新 m_team 名称，草稿列表取消该球队未关联标识"),

    ("", "在新标签页完成纳米匹配后，草稿页执行刷新、调整比赛顺序、添加比赛等重新渲染操作后再次打开弹窗，【匹配纳米数据源】按钮不再展示",
     "在新标签页已为该球队完成纳米匹配（nami_id_match_status 变为 1），但该球队尚未关联当前赛季",
     "1.切回草稿管理列表页，关闭原弹窗并执行刷新（或调整比赛顺序/添加比赛）\n2.再次点击该球队未关联标识打开弹窗",
     "弹窗不再展示【匹配纳米数据源】按钮"),
]


def generate():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "草稿管理列表页-第13轮冲刺"

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # 标题区
    ws.merge_cells("A1:G1")
    ws["A1"] = "时间：2026-10-27  姓名："
    ws["A1"].font = Font(bold=True, size=11)
    ws["A1"].border = border

    ws.merge_cells("A2:F2")
    ws["A2"] = "执行开始时间"
    ws["A2"].font = Font(bold=True, size=11)
    ws["A2"].border = border

    headers = ["序号", "冲刺", "测试项", "测试点", "前置条件", "执行步骤", "预期结果", "测试结果1", "测试结果2", "测试结果3"]
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = Font(bold=True, size=11)
        c.border = border
        c.alignment = wrap

    row = 4
    for idx, (item, point, pre, steps, expect) in enumerate(cases, 1):
        data = [idx, "十三", item, point, pre, steps, expect]
        for col, v in enumerate(data, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = border
            c.alignment = wrap
        for col in (8, 9, 10):
            ws.cell(row=row, column=col).border = border
        row += 1

    widths = {"A": 5, "B": 6.75, "C": 16, "D": 38, "E": 36, "F": 26, "G": 46, "H": 9, "I": 9, "J": 9}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    wb.save(OUT)
    print("Saved module test cases:", OUT, "Total cases:", len(cases))

    os.makedirs(os.path.dirname(GLOBAL_OUT), exist_ok=True)
    shutil.copy2(OUT, GLOBAL_OUT)
    print("Archived to global test-cases:", GLOBAL_OUT)


if __name__ == "__main__":
    generate()
