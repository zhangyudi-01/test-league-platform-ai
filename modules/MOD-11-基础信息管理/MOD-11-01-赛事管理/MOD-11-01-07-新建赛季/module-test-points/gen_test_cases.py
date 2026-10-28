# -*- coding: utf-8 -*-
"""按模板格式生成 MOD-11-01-07 新建赛季 第13轮冲刺（PBL-636）测试用例 Excel"""
import os
import shutil
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

BASE_DIR = r"D:\pythonProject\test-league-platform-ai"
MOD_DIR = os.path.join(BASE_DIR, r"modules\MOD-11-基础信息管理\MOD-11-01-赛事管理\MOD-11-01-07-新建赛季")
OUT_DIR = os.path.join(MOD_DIR, "module-test-cases")
os.makedirs(OUT_DIR, exist_ok=True)
OUT = os.path.join(OUT_DIR, "测试用例-新建赛季-第13轮冲刺.xlsx")
GLOBAL_OUT = os.path.join(BASE_DIR, "test-cases", "测试用例-新建赛季-第13轮冲刺.xlsx")

# 列: A序号 B冲刺 C测试项 D测试点 E前置条件 F执行步骤 G预期结果 H测试结果1 I测试结果2 J测试结果3
cases = [
    ("匹配后页面去向", "新建赛季点击\"创建并继续\"进入匹配页面，确认匹配后刷新并停留在当前匹配页面",
     "在新建赛季页填写合规赛季信息，点击\"创建并继续\"创建成功并进入该赛季的匹配页面，列表存在可匹配的纳米赛季及阶段",
     "1.勾选可匹配的赛季及阶段\n2.点击【匹配】\n3.在确认提示框点击确认",
     "完成赛季匹配，页面刷新并停留在当前匹配页面"),

    ("", "新建赛季后确认匹配，不再跳转至当前赛季的关联球队页面",
     "同用例 1，在新建赛季后进入的匹配页面完成勾选并点击【匹配】",
     "1.在确认提示框点击确认\n2.查看页面是否跳转",
     "停留在当前赛季匹配页面，不跳转至当前赛季的关联球队页面"),

    ("", "刷新后，本次匹配的赛季及阶段显示为已匹配状态",
     "同用例 1，确认匹配后页面已完成刷新",
     "1.查看当前赛季的匹配状态及所勾选的赛季/阶段状态",
     "顶部匹配状态显示为\"已匹配\"，本次匹配的赛季及阶段记录显示为\"已匹配\"且复选框不可取消勾选"),

    ("", "匹配确认提示框点击取消，不进行匹配，停留在当前页面",
     "新建赛季点击\"创建并继续\"进入匹配页面，已勾选可匹配赛季并点击【匹配】弹出确认提示框",
     "1.在确认提示框点击取消",
     "确认提示框关闭，不进行匹配，停留在当前匹配页面"),

    ("停留后操作", "停留在匹配页面后点击\"返回\"，返回赛季详情页",
     "同用例 1，新建赛季确认匹配后停留在当前匹配页面",
     "1.点击页面底部的【返回】按钮",
     "返回赛季详情页"),

    ("回归", "从赛季详情页进入匹配赛季页面完成匹配，页面行为与本次调整前一致",
     "在赛季详情页选择某未匹配的已有赛季，进入匹配赛季页面",
     "1.勾选可匹配赛季及阶段，点击【匹配】并确认",
     "完成赛季匹配，页面刷新并停留在当前匹配页面"),
]


def generate():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "新建赛季-第13轮冲刺"

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # 标题区
    ws.merge_cells("A1:G1")
    ws["A1"] = "时间：2026-10-27  姓名："
    ws["A1"].font = Font(bold=True, size=11)
    ws["A1"].border = border

    ws.merge_cells("A2:F2")
    ws["A2"] = "执行开始时间"
    ws["A2"].font = Font(bold=True, size=11)
    ws["A2"].border = border

    headers = ["序号", "冲刺", "测试项", "测试点", "前置条件", "执行步骤", "预期结果", "测试结果1", "测试结果2", "测试结果3"]
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = Font(bold=True, size=11)
        c.border = border
        c.alignment = wrap

    row = 4
    for idx, (item, point, pre, steps, expect) in enumerate(cases, 1):
        data = [idx, "十三", item, point, pre, steps, expect]
        for col, v in enumerate(data, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = border
            c.alignment = wrap
        for col in (8, 9, 10):
            ws.cell(row=row, column=col).border = border
        row += 1

    widths = {"A": 5, "B": 6.75, "C": 16, "D": 38, "E": 36, "F": 26, "G": 46, "H": 9, "I": 9, "J": 9}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    wb.save(OUT)
    print("Saved module test cases:", OUT, "Total cases:", len(cases))

    os.makedirs(os.path.dirname(GLOBAL_OUT), exist_ok=True)
    shutil.copy2(OUT, GLOBAL_OUT)
    print("Archived to global test-cases:", GLOBAL_OUT)


if __name__ == "__main__":
    generate()
